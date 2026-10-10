"""投递记录存储（records.json）：加锁、原子写入、每日备份、旧记录迁移、查重、统计、Excel 镜像。"""

import copy
import fcntl
import json
import logging
import os
import re
import shutil
import tempfile
import threading
import time
import uuid
from contextlib import contextmanager
from datetime import datetime, timedelta

from . import config
from .checks import PUBLIC_DOMAINS, safe_url, split_emails
from .llm import COMPANY_TYPES

log = logging.getLogger(__name__)
_lock = threading.RLock()          # 同一进程内（面板的多个请求 / 后台队列线程）
_flock = {"depth": 0, "fh": None}   # 跨进程（面板和剪贴板投递同时写）
STATUSES = ["草稿", "已投递", "笔试", "已电联", "面试中", "offer", "拒绝", "无回复", "放弃"]   # 笔试：含在线测评
# 阶段：存储词不改（vault 按「已投递」「offer」「已电联」统计），界面统一用显示词（4.4）
STAGE_LABEL = {"草稿": "准备中", "已投递": "已投递", "笔试": "笔试/测评", "已电联": "已电联", "面试中": "面试", "offer": "offer",
               "拒绝": "未通过", "无回复": "无回复", "放弃": "放弃"}
STAGE_ORDER = {"草稿": 0, "已投递": 1, "笔试": 2, "已电联": 3, "面试中": 4, "offer": 5}
TERMINAL = ("拒绝", "无回复", "放弃")
POSITION_LABELS = ["全职", "留用实习", "实习", "不明确"]
KEEP_BACKUPS = 40   # 旧的每日备份留几份（backup.py 接手以后不再用）


class RecordsCorrupt(Exception):
    pass


def now_str():
    return datetime.now().strftime("%Y-%m-%d %H:%M")


# ── 读写 ───────────────────────────────────────────────────

@contextmanager
def _locked():
    """线程锁 + 文件锁（可重入）：保证读-改-写期间没有别的线程或进程插进来写。"""
    with _lock:
        if _flock["depth"] == 0:
            fh = open(config.RECORDS_PATH.with_name(".records.lock"), "a")
            fcntl.flock(fh, fcntl.LOCK_EX)
            _flock["fh"] = fh
        _flock["depth"] += 1
        try:
            yield
        finally:
            _flock["depth"] -= 1
            if _flock["depth"] == 0:
                fcntl.flock(_flock["fh"], fcntl.LOCK_UN)
                _flock["fh"].close()
                _flock["fh"] = None


def load_all():
    """整个 records.json：{"schema", "records", "applications", "updated_at", …}。schema 2 的文件 applications 当成空列表。"""
    with _lock:
        if not config.RECORDS_PATH.exists():
            return {"schema": 3, "records": [], "applications": []}
        text = config.RECORDS_PATH.read_text(encoding="utf-8")
        try:
            data = json.loads(text)
        except json.JSONDecodeError as e:
            raise RecordsCorrupt(
                f"records.json 损坏（{e}），为防止覆盖历史记录已停止写入。备份在 {config.BACKUP_DIR}") from e
        if isinstance(data, list):
            data = {"records": data}
        records = data.get("records") if isinstance(data, dict) else None
        if not isinstance(records, list):
            raise RecordsCorrupt("records.json 格式不对（没有 records 列表），已停止写入。")
        apps = data.setdefault("applications", [])
        if not isinstance(apps, list):
            raise RecordsCorrupt("records.json 格式不对（applications 不是列表），已停止写入。")
        for r in records:
            _fill_defaults(r)
        return data


def load():
    """只要岗位（看板上的记录）。"""
    return load_all()["records"]


def _write(data):
    from . import backup
    backup.tick()
    payload = json.dumps(data, ensure_ascii=False, indent=2)
    fd, tmp = tempfile.mkstemp(dir=config.RECORDS_PATH.parent, prefix=".records-", suffix=".tmp")
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as f:
            f.write(payload)
            f.flush()
            os.fsync(f.fileno())
        os.replace(tmp, config.RECORDS_PATH)
    finally:
        if os.path.exists(tmp):
            os.unlink(tmp)


def save_all(data):
    """写回整个文件（schema 3：岗位和申请一起）。不认识的顶层键原样保留。"""
    with _locked():
        data = dict(data)
        data["schema"] = max(3, data.get("schema") if isinstance(data.get("schema"), int) else 3)
        data["updated_at"] = now_str()
        data.setdefault("applications", [])
        _write(data)
    export_excel_safe(data["records"], data["applications"])


def save(records):
    """只换岗位列表：records、updated_at、schema 之外的顶层数据（applications 等）原样保留。"""
    with _locked():
        data = {}
        if config.RECORDS_PATH.exists():
            try:
                cur = json.loads(config.RECORDS_PATH.read_text(encoding="utf-8"))
            except json.JSONDecodeError as e:
                raise RecordsCorrupt(f"records.json 损坏（{e}），为防止覆盖历史记录已停止写入。备份在 {config.BACKUP_DIR}") from e
            data = cur if isinstance(cur, dict) else {}
        schema = data.get("schema") if isinstance(data.get("schema"), int) and data.get("schema") > 2 else 2
        data.update(records=records, updated_at=now_str(), schema=schema)
        _write(data)
    export_excel_safe(records, data.get("applications"))


def mutate(fn):
    """在锁内读-改-写岗位列表。fn(records) 的返回值原样返回。fn 里不要再调 add/update（外层保存会盖掉内层写入）。"""
    with _locked():
        records = load()
        out = fn(records)
        save(records)
        return out


def mutate_all(fn):
    """在锁内读-改-写整个文件（岗位和申请一起，要么一起写成功、要么都不写）。fn(data) 的返回值原样返回。"""
    with _locked():
        data = load_all()
        out = fn(data)
        save_all(data)
        return out


# ── 字段 ───────────────────────────────────────────────────

NEW_FIELDS = {
    "campaign": "", "position_type": "", "deadline": "", "source_url": "", "apply_channel": "",
    "target_job": "", "language": "", "resume_version": "", "resume_sha": "", "attachments": [],
    "gmail_message_id": "", "gmail_thread_id": "", "gmail_draft_id": "", "send_mode": "",
    "reply_status": "", "reply_at": "", "reply_from": "", "reply_snippet": "", "reply_checked_at": "",
    "issues_at_send": [], "model": "", "cc_email": "", "notes": "", "focus_industry": "",
    "job_source": "", "job_post_date": "", "apply_url": "", "status_updated_at": "",
    "platform": "", "wangshen": {},
    "apply_account": "",     # 网申用哪个账号投的（注册网申账号的手机号 / 邮箱）
    "ws_submitted": "",      # 网申实际提交的内容：从网站上读回来的原文（证件号打码），不是面板生成的
    "site_status": "",       # 网站上显示的进度（如「笔试」「已进入测评环节」），和查的时间
    "site_status_at": "",
    "wangshen_unused": {},   # 面板生成过、但本人没用上的网申问答：挪到这里，不再当成投递内容显示
    "sent_ts": 0,            # 投递时刻的绝对时间戳（换时区也准；旧记录为 0，按 sent_at 本机时间算）
    "reply_locked": False,   # 看板里手动改过回复状态：查回复时不再覆盖
    # v3：岗位挂在申请下面（4.1、4.2）
    "app_id": "",            # 属于哪个申请
    "org": "",               # 实际机构或部门（某理财子公司、某基金子公司……）
    "choice_no": 0,          # 志愿序号（第几志愿；0 = 没有志愿这一说）
    "accept_adjust": "",     # 是否服从调剂（是 / 否 / 空）
    "exam_city": "",         # 笔试城市
    "jd_url": "", "jd_at": "", "jd_source": "",   # JD 从哪读的、什么时候、谁读的（网页脚本 / 助手照抄 / 投递页）
    "uploaded_files": [],    # 这个岗位传过的文件
    "history": [],           # 阶段变化：[{at, from, to, by, reason}]（不再写进备注）
    "reply_subject": "",     # 来信标题
    "reply_kind": "",        # 来信认出来是哪一类进展（笔试 / 测评 / AI 面 / 面试邀请、拒信、offer）
}


def _fill_defaults(r):
    for k, v in NEW_FIELDS.items():
        if k not in r or r[k] is None:
            r[k] = copy.deepcopy(v)
    if not r["campaign"]:
        r["campaign"] = config.OLD_CAMPAIGN if (r.get("sent_at") or "") < config.OLD_CAMPAIGN_BEFORE else config.CURRENT_CAMPAIGN
    if not r["position_type"] and r["campaign"] == config.OLD_CAMPAIGN:
        r["position_type"] = "实习"
    if not r["send_mode"]:
        r["send_mode"] = "发送" if r.get("to_email") else "未发邮件"
    if not r["resume_version"] and r["campaign"] == config.OLD_CAMPAIGN:
        r["resume_version"] = "旧版简历"
    return r


def migrate():
    """把旧记录补齐新字段并落盘（幂等）。返回补齐的条数。"""
    if not config.RECORDS_PATH.exists():
        return 0
    load()  # 文件坏了会抛 RecordsCorrupt（友好提示），不会覆盖
    raw = json.loads(config.RECORDS_PATH.read_text(encoding="utf-8"))
    raw_records = raw.get("records") if isinstance(raw, dict) else raw
    missing = sum(1 for r in raw_records if any(k not in r for k in NEW_FIELDS))
    if missing:
        mutate(lambda recs: None)  # load() 已补默认值，save 落盘
    return missing


def new_record(**fields):
    r = copy.deepcopy(NEW_FIELDS)
    r.update({
        "id": uuid.uuid4().hex[:8], "company_name": "", "company_type": "", "job_title": "",
        "job_location": "", "to_email": "", "subject": "", "email_body": "", "jd_text": "",
        "sent_at": now_str(), "sent_ts": time.time(), "status": "已投递", "source_type": "网页面板",
        "attach_report": False, "created_at": now_str(), "status_updated_at": now_str(),
        "campaign": config.CURRENT_CAMPAIGN,
    })
    r.update(fields)
    _clean_urls(r)
    return r


def _clean_urls(r):
    """链接只留 http(s)，防止 javascript: 之类的东西进看板。"""
    for k in ("apply_url", "source_url"):
        if k in r:
            r[k] = safe_url(r.get(k))


def add(record):
    """加一条岗位；它所属的申请没有就顺手建好（links 迁移之后旧代码新建的记录也有申请）。"""
    def _add(data):
        from . import apps
        data["records"].append(record)
        apps.ensure_in(data, record)
        return record["id"]
    return mutate_all(_add)


def get(record_id):
    return next((r for r in load() if r.get("id") == record_id), None)


def update(record_id, fields, by="面板"):
    """更新字段。状态变化记进这个岗位的 history 和它申请的时间线（不再往备注里写）。"""
    def _upd(data):
        for r in data["records"]:
            if r.get("id") != record_id:
                continue
            upd = dict(fields)
            new_status = upd.get("status")
            old_status = r.get("status", "")
            if new_status and new_status != old_status:
                stamp = now_str()
                upd["status_updated_at"] = stamp
                r.setdefault("history", []).append({"at": stamp, "from": old_status, "to": new_status, "by": by})
                app = next((a for a in data.get("applications", []) if a.get("id") == r.get("app_id")), None)
                if app is not None:
                    app.setdefault("timeline", []).append(
                        {"at": stamp, "kind": "阶段", "by": by,
                         "text": f"{r.get('job_title') or '岗位'}：{STAGE_LABEL.get(old_status, old_status)} → {STAGE_LABEL.get(new_status, new_status)}"})
                    app["updated_at"] = stamp
            r.update(upd)
            _clean_urls(r)
            return True
        return False
    return mutate_all(_upd)


def delete(record_id):
    def _del(recs):
        before = len(recs)
        recs[:] = [r for r in recs if r.get("id") != record_id]
        return len(recs) < before
    return mutate(_del)


# ── 查重 ───────────────────────────────────────────────────

_COMPANY_NOISE = re.compile(
    r"[（(].*?[)）]|股份有限公司|有限责任公司|有限公司|私募基金管理|基金管理|投资管理|资产管理|"
    r"私募|创业投资|投资|资本|创投|基金|集团|控股|capital|partners|ventures|fund|management|"
    r"investments?|group|holdings|[\s·\-_.,，]", re.I)


def norm_company(name):
    return _COMPANY_NOISE.sub("", (name or "").lower())


def _domains(emails):
    return {e.split("@")[1] for e in emails if "@" in e and e.split("@")[1] not in PUBLIC_DOMAINS}


def find_related(company_name, emails, records=None, exclude_id=None):
    records = load() if records is None else records
    emails = split_emails(emails)
    new_set, new_domains = set(emails), _domains(emails)
    key = norm_company(company_name)
    out = []
    for r in records:
        if r.get("id") == exclude_id:
            continue
        old = split_emails(r.get("to_email")) + split_emails(r.get("cc_email"))
        match = ""
        if new_set & set(old):
            match = "同一邮箱"
        elif new_domains & _domains(old):
            match = "同一邮箱域名"
        else:
            other = norm_company(r.get("company_name"))
            if key and other and len(key) >= 2 and len(other) >= 2 and (key in other or other in key):
                match = "机构名相近"
        if match:
            out.append({"id": r.get("id"), "sent_at": r.get("sent_at", ""), "company_name": r.get("company_name", ""),
                        "job_title": r.get("job_title", ""), "status": r.get("status", ""),
                        "campaign": r.get("campaign", ""), "position_type": r.get("position_type", ""),
                        "reply_status": r.get("reply_status", ""), "match": match})
    out.sort(key=lambda x: x["sent_at"], reverse=True)
    return out


# ── 归一化 & 统计 ──────────────────────────────────────────

def norm_company_type(t):
    t = (t or "").strip()
    if t in COMPANY_TYPES:
        return t
    tl = t.lower()
    if not t:
        return "其他"
    if any(k in tl for k in ("fa", "投行", "券商", "并购顾问")):
        return "券商/投行/FA"
    if any(k in tl for k in ("国资", "政府", "引导")):
        return "国资/政府引导基金"
    if any(k in tl for k in ("险资", "银行", "资管")):
        return "险资/银行系/资管"
    if any(k in tl for k in ("战略", "cvc", "产业", "战投")):
        return "产业资本/CVC/战投"
    if any(k in tl for k in ("大厂", "互联网", "科技")):
        return "企业/大厂"
    if "vc" in tl and "pe" in tl:
        return "双币VC/PE"
    if "vc" in tl or "风险投资" in tl or "创投" in tl or "早期" in tl:
        return "美元VC" if "美元" in tl else "人民币VC"
    if "pe" in tl or "私募" in tl or "股权" in tl or "并购" in tl:
        return "PE/并购基金"
    return "其他"


def norm_city(loc):
    loc = (loc or "").strip()
    if not loc or loc.startswith("[") or loc in ("未注明", "不限", "不明确"):
        return ""
    for city in ("上海", "北京", "深圳", "杭州", "香港", "广州", "苏州", "南京", "成都", "新加坡"):
        if city in loc:
            return city
    if "线上" in loc or "远程" in loc:
        return "远程"
    return loc.split("/")[0].split("·")[0].strip()[:4] if loc else ""


_INDUSTRY_MAP = {"前沿科技": "硬科技", "科技": "硬科技", "云计算": "AI", "算力": "AI", "物理AI": "AI",
                 "AI基础设施": "AI", "具身智能": "AI", "集成电路": "半导体", "芯片": "半导体",
                 "EV产业链": "新能源", "食品饮料": "消费", "消费出海": "消费", "工业": "先进制造",
                 "制造": "先进制造", "全周期": "", "综合": "", "TMT/消费": "TMT"}


def norm_industries(value):
    out = []
    for part in re.split(r"[/、,，]", value or ""):
        p = _INDUSTRY_MAP.get(part.strip(), part.strip())
        if p and p not in out:
            out.append(p)
    return out


def filter_campaign(records, campaign):
    if not campaign or campaign == "全部":
        return records
    return [r for r in records if r.get("campaign") == campaign]


def stats(records):
    now = datetime.now()
    week_ago = now - timedelta(days=7)

    def _dt(r):
        try:
            return datetime.strptime((r.get("sent_at") or "")[:16], "%Y-%m-%d %H:%M")
        except ValueError:
            return None

    out = {"total": len(records), "week_new": 0, "followup": [], "status_dist": {}, "type_dist": {},
           "loc_dist": {}, "industry_dist": {}, "position_dist": {}, "daily": {}, "daily_detail": {},
           "replied": 0, "bounced": []}
    out["companies"] = len({norm_company(r.get("company_name")) for r in records if r.get("company_name")})
    for r in records:
        dt = _dt(r)
        status = r.get("status") or "已投递"
        out["status_dist"][status] = out["status_dist"].get(status, 0) + 1
        ctype = norm_company_type(r.get("company_type"))
        out["type_dist"][ctype] = out["type_dist"].get(ctype, 0) + 1
        city = norm_city(r.get("job_location"))
        if city:
            out["loc_dist"][city] = out["loc_dist"].get(city, 0) + 1
        for ind in norm_industries(r.get("focus_industry")):
            out["industry_dist"][ind] = out["industry_dist"].get(ind, 0) + 1
        pt = r.get("position_type") or "不明确"
        out["position_dist"][pt] = out["position_dist"].get(pt, 0) + 1
        if r.get("reply_status") == "有回复":
            out["replied"] += 1
        if r.get("reply_status") == "退信":
            out["bounced"].append(f"{r.get('company_name') or '未知'}（{r.get('to_email', '')}）")
        if dt:
            if dt >= week_ago:
                out["week_new"] += 1
            if status == "已投递" and dt < week_ago and r.get("reply_status") not in ("有回复", "退信"):
                out["followup"].append(r.get("company_name") or "未知")
            day = dt.strftime("%Y-%m-%d")  # 带年份排序（秋招会跨年），界面上只显示月/日
            out["daily"][day] = out["daily"].get(day, 0) + 1
            out["daily_detail"].setdefault(day, []).append(r.get("company_name", ""))
    return out


# ── Excel 镜像（方便直接用 Excel 看；失败不影响主流程）──────────

EXCEL_STATUS = {"ok": True, "error": "", "at": ""}
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _cell(v):
    if isinstance(v, bool):
        return "是" if v else ""
    if not isinstance(v, (str, int, float)):
        v = "" if v is None else str(v)
    return _ILLEGAL.sub("", v) if isinstance(v, str) else v


def _append(ws, row):
    """写一行；以 = 开头的文字按文字存，不让 Excel 当公式执行。"""
    ws.append([_cell(v) for v in row])
    for c in ws[ws.max_row]:
        if isinstance(c.value, str) and c.value.startswith("="):
            c.data_type = "s"


_PHONE = re.compile(r"(?<!\d)(1[3-9]\d)(\d{4})(\d{4})(?!\d)")


def _mask_id(text):
    from .apps import mask
    return mask(text)


def mask_phone(text):
    """手机号中间四位打码（桌面 Excel 里用）。"""
    return _PHONE.sub(r"\1****\3", text or "")


def _app_cols(r, apps_by_id):
    a = apps_by_id.get(r.get("app_id")) or {}
    nxt = a.get("next_step") or {}
    channel = a.get("channel") or r.get("apply_channel") or ("邮件" if r.get("to_email") else "网申")
    return {"_channel": channel, "_account": mask_phone(r.get("apply_account", "")),
            "_stage": STAGE_LABEL.get(r.get("status") or "已投递", r.get("status", "")),
            "_progress": r.get("site_status", ""),
            "_next": (nxt.get("text", "") + (f"（{nxt['due']}）" if nxt.get("due") else "")).strip(),
            "_link": f"{config.PANEL_BASE}/#app/{r['app_id']}" if r.get("app_id") else ""}


def export_excel(records, path=None, applications=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    path = path or config.EXCEL_MIRROR_PATH
    apps_by_id = {a.get("id"): a for a in applications or []}
    wb = Workbook()
    ws = wb.active
    ws.title = "投递记录"
    cols = [("投递时间", "sent_at", 17), ("批次", "campaign", 18), ("类型", "position_type", 9),
            ("机构", "company_name", 22), ("机构性质", "company_type", 14), ("岗位", "job_title", 26),
            ("地点", "job_location", 12), ("行业", "focus_industry", 14), ("状态", "status", 8),
            ("回复", "reply_status", 10), ("收件人", "to_email", 28), ("邮件标题", "subject", 40),
            ("简历版本", "resume_version", 14), ("附研究样本", "attach_report", 10),
            ("来源", "job_source", 18), ("截止", "deadline", 11), ("投递方式", "send_mode", 9),
            ("备注", "notes", 40),
            ("渠道", "_channel", 10), ("申请账号", "_account", 18), ("阶段", "_stage", 10), ("网站进度", "_progress", 22),
            ("下一步·截止", "_next", 26), ("面板链接", "_link", 30)]
    ws.append([c[0] for c in cols])
    for r in sorted(records, key=lambda x: x.get("sent_at") or "", reverse=True):
        extra = _app_cols(r, apps_by_id)
        row = []
        for _, key, _ in cols:
            v = extra[key] if key.startswith("_") else r.get(key, "")
            if key == "attach_report":
                v = "是" if v in (True, "是") else ""
            row.append(v)
        _append(ws, row)
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F4E79")
    for i, (_, _, width) in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font, cell.fill = head, fill
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    ws2 = wb.create_sheet("邮件与JD")
    ws2.append(["投递时间", "机构", "岗位", "邮件正文", "JD 原文"])
    for r in sorted(records, key=lambda x: x.get("sent_at") or "", reverse=True):
        _append(ws2, [r.get("sent_at", ""), r.get("company_name", ""), r.get("job_title", ""),
                      (r.get("email_body") or "")[:32000], (r.get("jd_text") or "")[:32000]])
    for col, width in zip("ABCDE", (17, 22, 26, 60, 90)):
        ws2.column_dimensions[col].width = width
    for row in ws2.iter_rows(min_row=2):
        for cell in row[3:]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws2.freeze_panes = "A2"

    ws3 = wb.create_sheet("网申实际提交")   # 每个网申申请一行：网站上真交上去的内容（从网站读回的原文）
    ws3.append(["首次投出", "机构", "岗位", "申请账号", "网站进度", "实际提交（从网站读回）"])
    by_app = {}
    for r in records:
        if r.get("app_id"):
            by_app.setdefault(r["app_id"], []).append(r)
    rows = []
    for a in applications or []:
        if "网申" not in (a.get("channel") or "") or a.get("deleted_at"):
            continue
        pos = by_app.get(a.get("id"), [])
        text = next((p.get("ws_submitted") for p in pos if p.get("ws_submitted")), "")
        acct = a.get("account") or {}
        rows.append([a.get("first_submitted_at", ""), a.get("company", ""), "、".join(p.get("job_title", "") for p in pos),
                     mask_phone(acct.get("login") or acct.get("form_phone") or next((p.get("apply_account") for p in pos if p.get("apply_account")), "")),
                     (a.get("site_progress") or {}).get("text") or next((p.get("site_status") for p in pos if p.get("site_status")), ""),
                     mask_phone(_mask_id(text))[:32000]])
    for row in sorted(rows, key=lambda x: x[0] or "", reverse=True):
        _append(ws3, row)
    for col, width in zip("ABCDEF", (17, 22, 40, 18, 22, 100)):
        ws3.column_dimensions[col].width = width
    for row in ws3.iter_rows(min_row=2):
        row[5].alignment = Alignment(wrap_text=True, vertical="top")
    ws3.freeze_panes = "A2"
    wb.save(path)
    return path


def export_excel_safe(records, applications=None):
    """Excel 正被打开等情况：跳过，下次保存时再同步；原因记在 EXCEL_STATUS 里给面板显示。"""
    try:
        if config.MATERIALS_DIR.exists():
            export_excel(records, applications=applications)
        EXCEL_STATUS.update(ok=True, error="", at=now_str())
    except Exception as e:
        log.warning("Excel 镜像同步失败：%s", e)
        EXCEL_STATUS.update(ok=False, error=f"{type(e).__name__}: {e}", at=now_str())
