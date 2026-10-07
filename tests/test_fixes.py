"""这一轮复查修掉的问题 + 网申功能的测试（不调 AI、不碰 Gmail、不发信）。"""

import json
import subprocess
import sys
import threading
import time

import pytest
from openpyxl import load_workbook

import app as panel
from jobapply import checks, config, gmail_client, jobqueue, pipeline, records

N = config.CANDIDATE_NAME
LOCAL = {"Host": "localhost:5001"}
XRW = {"X-Requested-With": "jobapply"}


def result(**kw):
    r = {"company_name": "某资本", "job_title": "投资实习生", "jd_language": "中文", "apply_channel": "邮箱",
         "to_emails": ["hr@abc-capital.com"], "cc_emails": [], "apply_url": "", "contact_in_jd": "",
         "jd_rules": {}, "resume_version": "双语", "resume_filename": f"{N}-简历.pdf", "attach_report": False,
         "email_subject": f"投资实习生申请 - {N}",
         "email_body": f"您好，\n\n我是{N}，硕士在读，两周内可到岗，每周5天。\n\n简历见附件。\n\n{N}\n"}
    r.update(kw)
    return r


def to_issues(jd, **kw):
    return [i for i in checks.run(result(**kw), jd) if i["field"] == "to"]


# ── 收件邮箱：必须和 JD 里的完全一致 ───────────────────────────

def test_email_substring_of_jd_email_is_error():
    # 以前用「包含」判断：r@abc-capital.com 是 hr@abc-capital.com 的一部分，会被放过
    issues = to_issues("简历发送至 hr@abc-capital.com", to_emails=["r@abc-capital.com"])
    assert any(i["level"] == "error" for i in issues)


def test_exact_email_passes_and_obfuscated_warns():
    assert to_issues("简历发送至 hr@abc-capital.com") == []
    for jd in ("简历发送至 hr[at]abc-capital.com", "简历发送至 hr#abc-capital.com", "简历发送至 hr（at）abc-capital.com"):
        lv = [i["level"] for i in to_issues(jd)]
        assert lv == ["warn"], (jd, lv)


def test_email_read_from_image_no_warning():
    """图片里的邮箱抓取时已自动双重核对：不再提示你去看图。"""
    jd = "某资本招聘，详见下图。\n\n" + checks.OCR_MARKER + "\n投递邮箱：hr@abc-capital.com"
    assert to_issues(jd) == []
    # 正文里本来就打出来的邮箱，不算图片识别的
    jd2 = "投递邮箱：hr@abc-capital.com\n\n" + checks.OCR_MARKER + "\n投递邮箱：hr@abc-capital.com"
    assert to_issues(jd2) == []


# ── 称呼 ───────────────────────────────────────────────────

@pytest.mark.parametrize("line,name", [
    ("尊敬的王总：", "王"), ("李老师好！", "李"), ("王经理，您好！", "王"), ("张女士您好，", "张"),
    ("Dear Ms. Chen,", "Chen"), ("Hi Kevin,", "Kevin"), ("您好，", ""), ("各位老师好，", ""), ("HR您好，", ""),
    ("Hello,", ""), ("Hi there,", ""), ("Dear Hiring Manager,", ""), ("老师您好！", ""),
])
def test_greeting_name(line, name):
    assert checks._greeting_name(line) == name


def test_greeting_without_jd_evidence_replaced():
    jd = "简历发送至 hr@abc-capital.com"
    r = result(email_body=f"尊敬的王总：\n\n正文足够长足够长足够长足够长足够长足够长足够长。\n\n{N}\n")
    fixes = checks.autofix(r, jd)
    assert r["email_body"].startswith("您好，") and fixes


# ── 占位符：JD 自己要求的方括号写法不算 ─────────────────────────

def test_bracket_required_by_jd_is_not_placeholder():
    jd = "邮件标题请注明：[实习申请]姓名-学校，简历发 hr@abc-capital.com"
    r = result(email_subject=f"[实习申请]{N}-某大学", jd_rules={"subject_format": "[实习申请]姓名-学校"})
    assert not [i for i in checks.run(r, jd) if i["field"] == "subject" and "占位" in i["msg"]]
    r2 = result(email_subject=f"[期望日薪]{N}")
    assert [i for i in checks.run(r2, jd) if i["field"] == "subject" and "占位" in i["msg"]]


# ── 链接只放行 http(s) ──────────────────────────────────────

@pytest.mark.parametrize("raw,clean", [
    ("https://jobs.feishu.cn/abc", "https://jobs.feishu.cn/abc"),
    ("www.mokahr.com/apply/x", "https://www.mokahr.com/apply/x"),
    ("javascript:alert(1)", ""), ("data:text/html,hi", ""), ("hr@abc.com", ""), ("扫码投递", ""),
    ("网申链接：https://a.com/x，截止 10 月底", "https://a.com/x"), ("", ""),
])
def test_safe_url(raw, clean):
    assert checks.safe_url(raw) == clean


def test_autofix_cleans_apply_url():
    r = result(apply_url="javascript:alert(1)")
    checks.autofix(r, "简历发送至 hr@abc-capital.com")
    assert r["apply_url"] == ""


# ── 记录存储 ───────────────────────────────────────────────

@pytest.fixture
def tmp_store(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "RECORDS_PATH", tmp_path / "records.json")
    monkeypatch.setattr(config, "BACKUP_DIR", tmp_path / "backups")
    monkeypatch.setattr(config, "MATERIALS_DIR", tmp_path)
    monkeypatch.setattr(config, "EXCEL_MIRROR_PATH", tmp_path / "mirror.xlsx")
    return tmp_path


def test_new_records_do_not_share_defaults():
    a, b = records.new_record(), records.new_record()
    a["wangshen"]["x"] = 1
    a["attachments"].append("y")
    assert b["wangshen"] == {} and b["attachments"] == []
    assert records.NEW_FIELDS["wangshen"] == {}


def test_record_urls_cleaned(tmp_store):
    rid = records.add(records.new_record(company_name="A", apply_url="javascript:alert(1)", source_url="https://mp.weixin.qq.com/s/x"))
    rec = records.get(rid)
    assert rec["apply_url"] == "" and rec["source_url"] == "https://mp.weixin.qq.com/s/x"


def test_excel_survives_control_chars_and_formulas(tmp_store):
    records.add(records.new_record(company_name="=HYPERLINK(\"http://evil\",\"点我\")", subject="标题\x07带响铃",
                                   email_body="正文\x00\x1f", jd_text="=1+1"))
    assert records.EXCEL_STATUS["ok"], records.EXCEL_STATUS
    wb = load_workbook(tmp_store / "mirror.xlsx")
    ws = wb["投递记录"]
    company = ws.cell(row=2, column=4)
    assert company.data_type == "s" and company.value.startswith("=HYPERLINK")
    assert "\x07" not in ws.cell(row=2, column=12).value


def test_excel_failure_is_reported_not_raised(tmp_store, monkeypatch):
    def boom(*a, **k):
        raise PermissionError("文件被占用")
    monkeypatch.setattr(records, "export_excel", boom)
    records.add(records.new_record(company_name="A"))
    assert not records.EXCEL_STATUS["ok"] and "被占用" in records.EXCEL_STATUS["error"]
    monkeypatch.undo()


def test_file_lock_is_reentrant(tmp_store):
    # 锁里再调一次写（同一线程）：不能卡死
    t = threading.Thread(target=lambda: records.mutate(lambda recs: records.add(records.new_record(company_name="嵌套"))))
    t.start()
    t.join(timeout=10)
    assert not t.is_alive()


def test_file_lock_waits_for_other_process(tmp_store):
    # 剪贴板投递（另一个进程）正在写记录时，面板要等它写完，不能同时写
    lock = config.RECORDS_PATH.with_name(".records.lock")
    code = (f"import fcntl,time; f=open({str(lock)!r},'a'); fcntl.flock(f, fcntl.LOCK_EX); "
            "print('locked', flush=True); time.sleep(1.0)")
    p = subprocess.Popen([sys.executable, "-c", code], stdout=subprocess.PIPE, text=True)
    try:
        assert p.stdout.readline().strip() == "locked"
        t0 = time.time()
        records.add(records.new_record(company_name="等锁"))
        assert time.time() - t0 >= 0.5
    finally:
        p.wait(timeout=10)
    assert records.load()[0]["company_name"] == "等锁"


def test_wangshen_merged_into_sent_record(tmp_store):
    rid = records.add(records.new_record(company_name="某资本", send_mode="发送"))
    kit = {"platform": "飞书", "self_intro": "我是……", "junk": "不该存"}
    out = pipeline.record_web_application(result(), "JD", kit=kit, upload_version="中文", record_id=rid)
    assert out["merged"] and out["record_id"] == rid and len(records.load()) == 1
    rec = records.get(rid)
    assert rec["platform"] == "飞书" and "junk" not in rec["wangshen"] and "已同时网申" in rec["notes"]


def test_web_application_recorded(tmp_store):
    out = pipeline.record_web_application(result(to_emails=[], apply_channel="网申/链接"), "JD",
                                          kit={"platform": "Moka"}, upload_version="英文")
    rec = records.get(out["record_id"])
    assert rec["send_mode"] == "未发邮件" and rec["resume_version"] == "网申上传（英文）" and rec["platform"] == "Moka"


# ── 批量队列：防重复发送 / 授权失效 / 中断恢复 ────────────────────

@pytest.fixture
def q(tmp_path, monkeypatch):
    monkeypatch.setattr(jobqueue, "QUEUE_PATH", tmp_path / "queue.json")
    monkeypatch.setattr(jobqueue, "_maybe_notify", lambda: None)
    monkeypatch.setattr(jobqueue, "_submit", lambda item_id: None)  # 不跑后台分析
    return jobqueue


def add_ready(q, **kw):
    it = q._new_item("text", "JD 文字" * 20)
    it.update(status="待审核", jd_text="简历发送至 hr@abc-capital.com",
              analysis={"result": result(**kw), "issues": [], "fixes": [], "meta": {}})
    items = q._load()
    items.append(it)
    q._save(items)
    return it["id"]


def test_split_input_title_plus_link():
    share = "【招聘】某资本2027届投资实习生 https://mp.weixin.qq.com/s/AbC123"
    assert jobqueue.split_input(share) == [("url", "https://mp.weixin.qq.com/s/AbC123")]
    jd = "岗位职责：负责行业研究。" * 10 + "\n网申链接：https://jobs.feishu.cn/x"
    assert jobqueue.split_input(jd) == [("text", jd)]


def test_claim_only_once(q):
    iid = add_ready(q)
    assert q.claim(iid) and not q.claim(iid)
    assert q._get(iid)["status"] == "发送中"
    q.release(iid, error="网络断了")
    it = q._get(iid)
    assert it["status"] == "待审核" and it["error"] == "网络断了"


def test_save_edits_normalizes_and_respects_claim(q):
    iid = add_ready(q)
    assert q.save_edits(iid, {"to_emails": "a@x.com; b@y.com"}, issues=[{"level": "warn", "msg": "m"}])
    it = q._get(iid)
    assert it["edited"]["to_emails"] == ["a@x.com", "b@y.com"] and it["live_issues"][0]["msg"] == "m"
    assert q.save_edits(iid, wangshen={"self_intro": "改过"}) and q._get(iid)["wangshen"]["self_intro"] == "改过"
    q.claim(iid)
    assert not q.save_edits(iid, {"to_emails": "c@z.com"})


def test_interrupted_send_needs_manual_check(q):
    iid = add_ready(q)
    q.claim(iid)
    q.resume_pending()
    it = q._get(iid)
    assert it["status"] == "需处理" and "Gmail" in it["error"]


def test_process_ready_spacing_and_claims(q, monkeypatch):
    ids = [add_ready(q), add_ready(q)]
    monkeypatch.setattr(pipeline, "review", lambda *a, **k: {"issues": []})
    sent, sleeps = [], []
    monkeypatch.setattr(pipeline, "deliver", lambda r, jd, **kw: sent.append(kw["mode"]) or {"ok": True, "mode": "发送", "record_id": "r"})
    monkeypatch.setattr(jobqueue.time, "sleep", lambda s: sleeps.append(s))
    out = q.process_ready("send")
    assert len(out["done"]) == 2 and sent == ["send", "send"] and sleeps == [jobqueue.SEND_GAP_SECONDS]
    assert {q._get(i)["status"] for i in ids} == {"已发送"}


def test_process_ready_skips_mail_plus_wangshen(q, monkeypatch):
    add_ready(q, apply_channel="邮箱+网申")
    monkeypatch.setattr(pipeline, "review", lambda *a, **k: {"issues": []})
    monkeypatch.setattr(pipeline, "deliver", lambda *a, **k: pytest.fail("不该发送"))
    out = q.process_ready("send")
    assert out["done"] == [] and "网申" in out["skipped"][0]


def test_process_ready_stops_on_expired_gmail(q, monkeypatch):
    iid = add_ready(q)
    monkeypatch.setattr(pipeline, "review", lambda *a, **k: {"issues": []})

    def expired(*a, **k):
        raise gmail_client.GmailAuthError("授权过期")
    monkeypatch.setattr(pipeline, "deliver", expired)
    with pytest.raises(gmail_client.GmailAuthError):
        q.process_ready("send")
    assert q._get(iid)["status"] == "待审核"


def test_process_ready_one_at_a_time(q):
    assert jobqueue._ready_lock.acquire(blocking=False)
    try:
        with pytest.raises(jobqueue.Busy):
            q.process_ready("send")
    finally:
        jobqueue._ready_lock.release()


def test_retry_refuses_sending_item(q):
    iid = add_ready(q)
    q.claim(iid)
    assert q.retry(iid) is False


# ── 面板接口 ───────────────────────────────────────────────

@pytest.fixture
def client(q, tmp_store):
    panel.app.config["TESTING"] = True
    return panel.app.test_client()


def test_send_refuses_claimed_queue_item(client, monkeypatch):
    iid = add_ready(jobqueue)
    jobqueue.claim(iid)
    monkeypatch.setattr(pipeline, "deliver", lambda *a, **k: pytest.fail("不该发送"))
    r = client.post("/api/send", json={"result": result(), "jd_text": "x", "queue_id": iid, "mode": "send"}, headers={**LOCAL, **XRW})
    assert r.status_code == 409 and r.get_json()["queue_taken"]


def test_send_failure_releases_claim(client, monkeypatch):
    iid = add_ready(jobqueue)

    def blocked(*a, **k):
        raise pipeline.Blocked([{"level": "error", "field": "to", "msg": "邮箱不对"}])
    monkeypatch.setattr(pipeline, "deliver", blocked)
    r = client.post("/api/send", json={"result": result(), "jd_text": "x", "queue_id": iid, "mode": "send"}, headers={**LOCAL, **XRW})
    assert r.status_code == 409 and jobqueue._get(iid)["status"] == "待审核"


def test_send_success_marks_queue_done(client, monkeypatch):
    iid = add_ready(jobqueue)
    monkeypatch.setattr(pipeline, "deliver", lambda *a, **k: {"ok": True, "mode": "草稿", "record_id": "r1", "attachments": []})
    r = client.post("/api/send", json={"result": result(), "jd_text": "x", "queue_id": iid, "mode": "draft"}, headers={**LOCAL, **XRW})
    assert r.status_code == 200 and jobqueue._get(iid)["status"] == "已存草稿"


def test_record_endpoint_and_merge(client):
    iid = add_ready(jobqueue, to_emails=[], apply_channel="网申/链接")
    body = {"result": result(to_emails=[], apply_channel="网申/链接"), "jd_text": "JD", "queue_id": iid,
            "kit": {"platform": "北森", "self_intro": "我是"}, "upload_version": "中文"}
    d = client.post("/api/record", json=body, headers={**LOCAL, **XRW}).get_json()
    assert jobqueue._get(iid)["status"] == "已记录" and records.get(d["record_id"])["platform"] == "北森"
    # 再点一次：已经记过了，不重复记
    assert client.post("/api/record", json=body, headers={**LOCAL, **XRW}).status_code == 409
    # 邮箱+网申：邮件已发（队列条目已是「已发送」），带 record_id 补记到原记录上
    iid2 = add_ready(jobqueue, apply_channel="邮箱+网申")
    rid = records.add(records.new_record(company_name="某资本", send_mode="发送"))
    jobqueue.mark_done(iid2, "已发送", rid)
    d2 = client.post("/api/record", json={**body, "queue_id": iid2, "record_id": rid}, headers={**LOCAL, **XRW}).get_json()
    assert d2["merged"] and len(records.load()) == 2


def test_queue_put_saves_issues_and_kit(client):
    iid = add_ready(jobqueue)
    r = client.put(f"/api/queue/{iid}", json={"edited": result(to_emails="a@x.com"), "issues": [{"level": "info", "msg": "i"}],
                                              "wangshen": {"why_this_role": "因为"}}, headers={**LOCAL, **XRW})
    assert r.get_json()["ok"]
    it = jobqueue._get(iid)
    assert it["edited"]["to_emails"] == ["a@x.com"] and it["wangshen"]["why_this_role"] == "因为" and it["live_issues"]


def test_kit_endpoint(client):
    d = client.get("/api/kit", headers=LOCAL).get_json()
    assert d["基本信息"] and d["实习经历"]


def test_wangshen_endpoint(client, monkeypatch):
    assert client.post("/api/wangshen", json={"jd_text": "短"}, headers={**LOCAL, **XRW}).status_code == 400
    iid = add_ready(jobqueue, to_emails=[])
    monkeypatch.setattr(pipeline, "wangshen", lambda jd, **kw: ({"platform": "飞书", "self_intro": "我是"}, {"backend": "测试"}))
    d = client.post("/api/wangshen", json={"jd_text": "岗位职责" * 20, "result": result(), "queue_id": iid},
                    headers={**LOCAL, **XRW}).get_json()
    assert d["kit"]["platform"] == "飞书" and jobqueue._get(iid)["wangshen"]["platform"] == "飞书"


def test_stats_reports_excel_status(client):
    assert "excel" in client.get("/api/stats?campaign=全部", headers=LOCAL).get_json()


# ── Gmail 回复检查 ─────────────────────────────────────────

@pytest.mark.parametrize("headers,kind", [
    ({"subject": "感谢您投递某资本，诚邀您参加笔试"}, "有回复"),
    ({"subject": "【Moka】投递成功通知"}, "自动回复"),
    ({"subject": "Re: 实习申请", "auto-submitted": "auto-replied"}, "自动回复"),
    ({"subject": "Re: 实习申请"}, "有回复"),
    ({"subject": "Interview invitation", "auto-submitted": "auto-generated"}, "有回复"),
    ({"subject": "Undeliverable: 实习申请", "from": "postmaster@x.com"}, "退信"),
])
def test_reply_kind(headers, kind):
    assert gmail_client._reply_kind(headers) == kind


def test_company_terms():
    assert gmail_client._company_terms("红杉中国（HongShan）") == (["红杉中国", "HongShan"], [])
    assert gmail_client._company_terms("高瓴资本") == (["高瓴资本"], ["高瓴"])   # 短名字只在发件人里找
    assert gmail_client._company_terms("某资本（北京）") == (["某资本"], [])      # 括号里的城市不当搜索词
    assert gmail_client._company_terms("IDG资本") == (["IDG资本"], ["idg"])
    assert gmail_client._company_terms("") == ([], [])


def fake_records():
    return [{"id": str(i), "send_mode": "发送", "to_email": f"hr{i}@x.com", "company_name": f"机构{i}",
             "sent_at": "2026-10-01 10:00"} for i in range(8)]


def test_check_replies_skips_single_failures(monkeypatch):
    monkeypatch.setattr(gmail_client, "get_service", lambda: object())
    calls = []

    def one(svc, r, *a):
        calls.append(r["id"])
        if r["id"] == "2":
            raise TimeoutError("超时")
        return {"reply_checked_at": "now"}
    monkeypatch.setattr(gmail_client, "_check_one", one)
    out = gmail_client.check_replies(fake_records())
    assert len(calls) == 8 and "2" not in out and len(out) == 7


def test_check_replies_stops_after_repeated_failures_but_keeps_results(monkeypatch):
    monkeypatch.setattr(gmail_client, "get_service", lambda: object())

    def flaky(svc, r, *a):
        if r["id"] == "0":
            return {"reply_status": "有回复"}
        raise ConnectionError("断网")
    monkeypatch.setattr(gmail_client, "_check_one", flaky)
    logs = []
    out = gmail_client.check_replies(fake_records(), progress=logs.append)
    assert out == {"0": {"reply_status": "有回复"}} and any("已停下" in l for l in logs)


def test_check_replies_auth_expired(monkeypatch):
    monkeypatch.setattr(gmail_client, "get_service", lambda: object())

    def expired(svc, r, *a):
        raise gmail_client.RefreshError("invalid_grant")
    monkeypatch.setattr(gmail_client, "_check_one", expired)
    with pytest.raises(gmail_client.GmailAuthError):
        gmail_client.check_replies(fake_records())


def test_web_applications_are_checked_too(monkeypatch):
    monkeypatch.setattr(gmail_client, "get_service", lambda: object())
    seen = []
    monkeypatch.setattr(gmail_client, "_check_one", lambda svc, r, *a: seen.append(r["id"]) or {})
    recs = [{"id": "w", "send_mode": "未发邮件", "company_name": "某资本", "sent_at": "2026-10-01 10:00"},
            {"id": "x", "send_mode": "未发邮件", "company_name": "", "sent_at": "2026-10-01 10:00"}]
    gmail_client.check_replies(recs)
    assert seen == ["w"]


# ── 剪贴板模式：有任何「注意」都要停下来问，不能倒计时自动发 ─────────

@pytest.fixture
def clip(monkeypatch):
    import apply as clipboard
    monkeypatch.setattr(clipboard, "ensure_gmail", lambda: True)
    monkeypatch.setattr(clipboard, "notify", lambda *a, **k: None)
    monkeypatch.setattr(clipboard.time, "sleep", lambda s: None)
    monkeypatch.setattr(clipboard.pyperclip, "paste", lambda: "岗位职责：投资研究。" * 10 + " 简历发送至 hr@abc-capital.com")
    monkeypatch.setattr(sys, "argv", ["apply.py"])
    return clipboard


def fake_analysis(issues):
    return {"ok": True, "result": result(position_type="实习", company_type="人民币VC", job_location="上海",
                                         resume_filename_en="", report_filename=""),
            "issues": issues, "fixes": [], "meta": {"model": "m", "seconds": 1, "backend": "会员额度"}}


def test_clipboard_stops_on_any_warning(clip, monkeypatch):
    monkeypatch.setattr(pipeline, "analyze", lambda *a, **k: fake_analysis([{"level": "warn", "field": "body", "msg": "正文有感叹号。"}]))
    asked = []
    monkeypatch.setattr(clip, "ask", lambda prompt: asked.append(prompt) or "")
    monkeypatch.setattr(pipeline, "deliver", lambda *a, **k: pytest.fail("有「注意」时不该自动发"))
    assert clip.main() == 0 and asked


def test_clipboard_auto_sends_when_clean(clip, monkeypatch):
    monkeypatch.setattr(pipeline, "analyze", lambda *a, **k: fake_analysis([{"level": "info", "field": "time", "msg": "提示"}]))
    monkeypatch.setattr(clip, "ask", lambda prompt: pytest.fail("全过时不该问"))
    sent = []
    monkeypatch.setattr(pipeline, "deliver", lambda r, jd, **kw: sent.append(kw["mode"]) or {"ok": True, "mode": "发送", "attachments": ["a.pdf"], "record_id": "r"})
    assert clip.main() == 0 and sent == ["send"]
