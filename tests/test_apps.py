"""WP4 数据模型：申请＋岗位。护栏 R1（只能往前推）、R2（首次投出只写一次）、R6（颜色不重复、能释放）、R8（影响清单）；
「轮到谁」按优先级表逐行；软删除和恢复；schema 2 文件能读能写；保存后 applications 还在；钩子；快照；备份；通知；Excel。"""

import json
from datetime import datetime, timedelta

import pytest

from jobapply import apps, backup, config, notify, records, wstasks


@pytest.fixture
def store(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "RECORDS_PATH", tmp_path / "records.json")
    monkeypatch.setattr(config, "BACKUP_DIR", tmp_path / "backups")
    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    return tmp_path


def rec(**kw):
    return records.new_record(company_name=kw.pop("company_name", "S资本"), job_title=kw.pop("job_title", "分析师"), **kw)


# ── 存储：schema 2 / 3、钩子 ─────────────────────────────────

def test_schema2_file_reads_and_writes(store):
    config.RECORDS_PATH.write_text(json.dumps({"records": [rec(id="r1")], "schema": 2}, ensure_ascii=False), encoding="utf-8")
    data = records.load_all()
    assert data["applications"] == [] and len(data["records"]) == 1
    records.save(records.load())                                         # 老接口：不升级、不丢东西
    assert json.loads(config.RECORDS_PATH.read_text(encoding="utf-8"))["schema"] == 2
    records.mutate_all(lambda d: d["applications"].append(apps.blank("a1", "邮件", "S资本")))
    data = json.loads(config.RECORDS_PATH.read_text(encoding="utf-8"))
    assert data["schema"] == 3 and data["applications"][0]["id"] == "a1"
    records.update("r1", {"notes": "本人写的"})                            # 之后任何保存都留着申请
    assert json.loads(config.RECORDS_PATH.read_text(encoding="utf-8"))["applications"][0]["id"] == "a1"


def test_add_hook_creates_app(store):
    rid = records.add(rec(to_email="hr@s.com"))
    r = records.get(rid)
    a = apps.get(r["app_id"])
    assert a["channel"] == "邮件" and a["company"] == "S资本" and a["first_submitted_at"] == r["sent_at"]
    rid2 = records.add(rec(apply_channel="网申/链接", status="草稿", send_mode="未发邮件"))
    a2 = apps.get(records.get(rid2)["app_id"])
    assert a2["channel"] == "网申" and a2["web_phase"] == "待开始" and not a2["first_submitted_at"]
    rid3 = records.add(rec(app_id="fixedid001"))                          # 带了 app_id：按这个 id 建
    assert records.get(rid3)["app_id"] == "fixedid001" and apps.get("fixedid001")


def test_task_hook_uses_task_id_and_links_records(store, monkeypatch):
    t = wstasks.add("https://jobs.example.com/v3", "V资本", "实习生")
    a = apps.get(t["id"])
    assert a["channel"] == "网申" and a["web_phase"] == "待开始" and a["entry_url"].endswith("/v3") and a["color"] == t["color"]
    t = wstasks.set_status(t["id"], "已填待提交")                          # 网申草稿记录挂到同一个申请
    assert records.get(t["record_id"])["app_id"] == t["id"]
    assert len([x for x in apps.list_apps() if x["company"] == "V资本"]) == 1   # 没有多出一个自动建的申请


def test_email_plus_web_merges_auto_app(store):
    rid = records.add(rec(company_name="W资本", to_email="hr@w.com"))      # 邮件先发：自动建了邮件申请
    auto = records.get(rid)["app_id"]
    t = wstasks.add("https://jobs.example.com/w", "W资本", "分析师", email_record_id=rid)
    assert records.get(rid)["app_id"] == t["id"] and apps.get(t["id"])["channel"] == "邮件+网申"
    with pytest.raises(apps.NotFound):
        apps.get(auto)                                                    # 空的自动申请并进来了，不留空壳


# ── R1 只能往前推 / R2 首次投出只写一次 ───────────────────────

def test_stage_forward_only(store):
    rid = records.add(rec(status="已投递"))
    r = apps.advance(rid, "笔试", by="面板")
    assert r["status"] == "笔试" and r["history"][-1]["to"] == "笔试"
    with pytest.raises(ValueError):
        apps.advance(rid, "已投递", by="面板")                           # 系统不能往回改
    r = apps.advance(rid, "已投递", by="本人", manual=True, reason="点错了")
    a = apps.get(r["app_id"])
    assert r["status"] == "已投递" and any("往回改" in e["text"] for e in a["timeline"])
    apps.advance(rid, "拒绝", by="面板")                                  # 往结束推：可以
    with pytest.raises(ValueError):
        apps.advance(rid, "面试中", by="面板")                           # 结束了再改算往回
    with pytest.raises(ValueError):
        apps.advance(rid, "乱写", by="本人", manual=True)


def test_first_submitted_once(store):
    a = apps.create("网申", "T银行")
    r1 = records.add(rec(company_name="T银行", status="草稿", app_id=a["id"], job_title="管培生"))
    r2 = records.add(rec(company_name="T银行", status="笔试", app_id=a["id"], job_title="研究岗"))
    a = apps.mark_submitted(a["id"], at="2026-10-08 21:29", basis="本人点了我已提交", by="本人")
    assert a["first_submitted_at"] == "2026-10-08 21:29" and a["web_phase"] == "已提交"
    assert records.get(r1)["status"] == "已投递" and records.get(r1)["sent_at"] == "2026-10-08 21:29"
    assert records.get(r2)["status"] == "笔试"                             # 走得更远的不动
    a = apps.mark_submitted(a["id"], at="2026-10-08 22:27", basis="助手又报了一次", by="助手")
    assert a["first_submitted_at"] == "2026-10-08 21:29" and records.get(r1)["sent_at"] == "2026-10-08 21:29"


def test_terminal_web_phase_not_changed_by_assistant(store):
    a = apps.create("网申", "U资本")
    apps.mark_submitted(a["id"], basis="x", by="本人")
    assert apps.set_web_phase(a["id"], "等你", by="助手")["web_phase"] == "已提交"      # R3
    assert apps.set_need(a["id"], "登录", "登录一下")["web_phase"] == "已提交"
    assert apps.set_web_phase(a["id"], "待开始", by="本人")["web_phase"] == "待开始"   # 本人可以


# ── R6 颜色 ──────────────────────────────────────────────────

def test_color_unique_and_released(store):
    ids = [apps.create("网申", f"C{i}")["id"] for i in range(8)]
    colors = [apps.pick_color(i) for i in ids]
    assert sorted(colors) == list(range(8))                               # 同时可见的不撞色
    extra = apps.create("网申", "第九家")["id"]
    c9 = apps.pick_color(extra)
    assert apps.color_label(apps.get(extra), apps.list_apps()).endswith("2")   # 超过 8 家：写成「蓝2」这种
    apps.give_up(ids[0], "不合适", by="本人")
    assert apps.get(ids[0])["color"] is None                              # 结束就释放
    newcomer = apps.create("网申", "第十家")["id"]
    assert apps.pick_color(newcomer) == colors[0] or c9 == colors[0]


# ── R8 放弃 / 删除 / 恢复 ────────────────────────────────────

def test_give_up_and_soft_delete_with_impact(store):
    a = apps.create("网申", "G资本")
    r1 = records.add(rec(company_name="G资本", status="草稿", app_id=a["id"]))
    r2 = records.add(rec(company_name="G资本", status="已投递", app_id=a["id"]))
    impact = apps.give_up(a["id"], "", by="本人", dry_run=True)
    assert impact["not_submitted"] == [r1] and impact["submitted"] == [r2]
    with pytest.raises(ValueError):
        apps.give_up(a["id"], " ", by="本人")                            # 必须写原因
    apps.give_up(a["id"], "地点不合适", by="本人")
    assert records.get(r1)["status"] == "放弃" and records.get(r2)["status"] == "已投递"
    assert apps.get(a["id"])["web_phase"] == "已放弃"
    imp = apps.soft_delete(a["id"], by="本人")
    assert {p["id"] for p in imp["positions"]} == {r1, r2}
    assert a["id"] not in {x["id"] for x in apps.list_apps()} and records.get(r1).get("deleted_at")
    apps.restore(a["id"])
    assert a["id"] in {x["id"] for x in apps.list_apps()} and not records.get(r1).get("deleted_at")
    apps.soft_delete(a["id"], by="本人")
    assert apps.purge_deleted(now=datetime.now() + timedelta(days=8)) == [a["id"]]
    assert records.get(r1) is None


def test_merge_rename_account(store):
    a, b = apps.create("网申", "J基金"), apps.create("网申", "J银行")
    r = records.add(rec(company_name="J基金", app_id=a["id"]))
    apps.merge(a["id"], b["id"], by="本人")
    assert records.get(r)["app_id"] == b["id"] and records.get(r)["company_name"] == "J银行"
    apps.rename(b["id"], "中国J银行")
    assert records.get(r)["company_name"] == "中国J银行"
    apps.set_account(b["id"], login="（页面上没显示）")                    # 占位话不收
    assert apps.get(b["id"])["account"]["login"] == ""
    apps.set_account(b["id"], login="19900000000")
    assert records.get(r)["apply_account"] == "19900000000"


# ── 快照 ─────────────────────────────────────────────────────

def test_snapshot_masks_and_dedupes(store):
    a = apps.create("网申", "K资本")
    s1 = apps.add_snapshot(a["id"], "投递记录", "身份证 123456200001011234 已提交", url="https://jobs.example.com/x")
    s2 = apps.add_snapshot(a["id"], "投递记录", "身份证 123456200001011234 已提交")   # 一样的内容：不重复存
    assert s1["id"] == s2["id"] and len(apps.get(a["id"])["snapshots"]) == 1
    text = apps.snapshot_text(s1)
    assert "[证件号已隐去]" in text and "123456200001011234" not in text
    assert (config.DATA_DIR / s1["path"]).exists()


# ── 轮到谁（4.6）逐行 ─────────────────────────────────────────

NOW = datetime(2026, 10, 9, 12, 0)


def turn(app, **ctx):
    ctx.setdefault("positions", [])
    ctx.setdefault("now", NOW)
    return apps.view(app, ctx)


def base(**kw):
    a = apps.blank("v1", kw.pop("channel", "网申"), "测试资本")
    a.update(kw)
    return a


@pytest.mark.parametrize("app,ctx,expect", [
    (base(channel="邮件"), {"positions": [{"status": "拒绝"}, {"status": "offer"}]}, "已结束"),
    (base(channel="邮件"), {"positions": [{"status": "草稿"}], "mail": {"status": "待审核"}}, "轮到你"),
    (base(channel="邮件"), {"positions": [{"status": "草稿"}], "mail": {"status": "需处理"}}, "轮到你"),
    (base(channel="邮件"), {"positions": [{"status": "草稿"}], "mail": {"status": "需处理", "send_uncertain": True}}, "轮到你"),
    (base(channel="邮件"), {"positions": [{"status": "草稿"}], "mail": {"status": "已定时", "send_at": "2026-10-09 08:00"}}, "轮到你"),
    (base(channel="邮件", reply={"status": "退信"}), {"positions": [{"status": "已投递"}]}, "轮到你"),
    (base(web_phase="等你", need={"kind": "登录", "text": "登录某银行"}), {"positions": [{"status": "草稿"}]}, "轮到你"),
    (base(web_phase="待你提交"), {"positions": [{"status": "草稿"}]}, "轮到你"),
    (base(web_phase="在填", questions=[{"q": "高中？"}]), {"positions": [{"status": "草稿"}], "agent": {"working": True}}, "轮到你"),
    (base(web_phase="已提交", readback={"state": "等你登录"}), {"positions": [{"status": "已投递"}]}, "轮到你"),
    (base(web_phase="已提交", readback={"state": "不完整", "attempts": 2, "missing": ["JD"]}), {"positions": [{"status": "已投递"}]}, "轮到你"),
    (base(web_phase="已提交", next_step={"text": "测评", "due": "2026-10-10 15:34"}), {"positions": [{"status": "已投递"}]}, "轮到你"),
    (base(web_phase="待开始"), {"positions": []}, "轮到你"),
    (base(web_phase="在填"), {"positions": [{"status": "草稿"}], "agent": {"working": False, "alive": False}}, "停了"),
    (base(web_phase="已提交", readback={"state": "读回中"}), {"positions": [{"status": "已投递"}], "agent": {}}, "停了"),
    (base(web_phase="在填"), {"positions": [{"status": "草稿"}], "agent": {"working": True}}, "面板在做"),
    (base(web_phase="在填"), {"positions": [{"status": "草稿"}], "agent": {"queued_why": "满了"}}, "面板在做"),
    (base(channel="邮件"), {"positions": [{"status": "草稿"}], "mail": {"status": "已定时", "send_at": "2026-10-09 18:00"}}, "面板在做"),
    (base(channel="邮件"), {"positions": [{"status": "草稿"}], "mail": {"status": "排队中"}}, "面板在做"),
    (base(web_phase="已提交", readback={"state": "完成"}), {"positions": [{"status": "已投递"}]}, "等对方"),
    (base(channel="邮件"), {"positions": [{"status": "已投递"}]}, "等对方"),
])
def test_turn_priority_table(app, ctx, expect):
    v = turn(app, **ctx)
    assert v["turn"] == expect, v
    assert v["turn_color"] == apps.TURN_COLOR[expect]


def test_turn_order_and_labels():
    v = turn(base(web_phase="等你", need={"kind": "登录", "text": "在蓝框网页里登录"}), positions=[{"status": "草稿"}],
             agent={"working": True})
    assert v["turn"] == "轮到你" and v["todo"] == "在蓝框网页里登录" and v["button"]["label"] == "我登录好了"
    v = turn(base(web_phase="在填"), positions=[{"status": "草稿"}], agent={"stopped_reason": "本人点了停下"})
    assert v["turn"] == "停了" and v["why"] == "本人点了停下" and v["button"]["label"] == "让助手接着做"
    v = turn(base(channel="邮件"), positions=[{"status": "面试中"}, {"status": "拒绝"}])
    assert v["stage"] == "面试中" and v["stage_label"] == "面试"
    assert apps.stage_of([{"status": "拒绝"}, {"status": "无回复"}]) == "拒绝"


# ── 通知 ─────────────────────────────────────────────────────

def test_notify_dedupes_and_quiet_hours(monkeypatch):
    sent = []
    monkeypatch.setattr(notify, "_mac", lambda t, x: sent.append(t))
    monkeypatch.setattr(notify, "_label", lambda app_id: "蓝·某银行" if app_id else "")
    notify._recent.clear()
    day = datetime(2026, 10, 9, 14, 0)
    assert notify.send("轮到你", "登录一下", app_id="a1", kind="turn", now=day)
    assert not notify.send("轮到你", "登录一下", app_id="a1", kind="turn", now=day + timedelta(minutes=5))   # 10 分钟内不重复
    assert notify.send("轮到你", "登录一下", app_id="a1", kind="turn", now=day + timedelta(minutes=11))
    night = datetime(2026, 10, 10, 2, 0)
    assert not notify.send("截止", "明天交", app_id="a2", kind="due", due="2026-10-10 18:00", now=night)    # 夜里静默
    assert notify.send("截止", "马上到期", app_id="a3", kind="due", due="2026-10-10 04:00", now=night)      # 3 小时内到期照发
    assert sent[0].startswith("蓝·某银行｜") and notify.events_since(0)[-1]["title"].startswith("蓝·某银行")


# ── 备份 ─────────────────────────────────────────────────────

def test_backup_hourly_daily_and_manual(store, monkeypatch):
    monkeypatch.setattr(backup, "HOURLY_KEEP", 3)
    records.add(rec())
    t0 = datetime(2026, 10, 20, 8, 0)                                    # 别跟上面保存时按真实时间备的那份撞日子
    made = backup.tick(t0)
    assert any(p.parent.name == "hourly" for p in made) and any(p.parent.name == "daily" for p in made)
    assert backup.tick(t0 + timedelta(minutes=10)) == []                 # 同一个小时不重复备份
    for h in range(1, 6):
        backup.tick(t0 + timedelta(hours=h))
    assert len(list((config.BACKUP_DIR / "hourly").iterdir())) == 3       # 只留最近几份
    daily = config.BACKUP_DIR / "daily" / "20261020"
    assert (daily / "records.json").exists() and (daily / "chats_snapshots.tar.gz").exists()
    m = backup.backup_now("迁移前 links")
    assert (m / "records.json").exists() and "links" in m.name


# ── 查重带回复状态、Excel 新列 ─────────────────────────────────

def test_find_related_has_reply_status(store):
    records.add(rec(company_name="R资本", to_email="hr@r.com", reply_status="退信"))
    rel = records.find_related("R资本", ["hr@r.com"])
    assert rel[0]["reply_status"] == "退信"


def test_excel_has_new_columns_and_web_sheet(store, tmp_path):
    from openpyxl import load_workbook
    t = wstasks.add("https://jobs.example.com/e", "E银行", "管培生")
    wstasks.update(t["id"], account="13800000000")
    wstasks.set_status(t["id"], "已提交")
    rid = records.get(wstasks.get(t["id"])["record_id"])["id"]
    records.update(rid, {"ws_submitted": "实际提交的简历：联系电话 13800000000", "site_status": "笔试"})
    data = records.load_all()
    path = records.export_excel(data["records"], path=tmp_path / "x.xlsx", applications=data["applications"])
    wb = load_workbook(path)
    head = [c.value for c in wb["投递记录"][1]]
    for col in ("渠道", "申请账号", "阶段", "网站进度", "下一步·截止", "面板链接"):
        assert col in head
    row = [c.value for c in wb["投递记录"][2]]
    assert row[head.index("申请账号")] == "138****0000" and row[head.index("网站进度")] == "笔试"
    assert row[head.index("面板链接")].endswith(f"#app/{t['id']}")
    web = wb["网申实际提交"]
    assert web.max_row == 2 and "138****0000" in web.cell(2, 6).value and "13800000000" not in web.cell(2, 6).value


def test_email_plus_web_email_record_joins_task_card(store):
    """「邮箱+网申」：邮件那条记录不另建一张「网申」卡；和网申待办关联后并进同一张卡（渠道邮件+网申）。"""
    t = wstasks.add("https://jobs.example.com/ew", "EW资本", "分析师")
    wstasks.update(t["id"], queue_id="q-ew")
    rid = records.add(records.new_record(company_name="EW资本", job_title="分析师", to_email="hr@ew.com", apply_channel="邮箱+网申",
                                         status="草稿", send_mode="草稿"))
    first = records.get(rid)["app_id"]
    assert apps.get(first)["channel"] == "邮件"                                  # 不是「网申」，也不会挂个「让助手填」
    wstasks.link_email_record("q-ew", rid)
    assert records.get(rid)["app_id"] == t["id"] and apps.get(t["id"])["channel"] == "邮件+网申"
    with pytest.raises(apps.NotFound):
        apps.get(first)                                                         # 自动建的那张空卡并掉了
